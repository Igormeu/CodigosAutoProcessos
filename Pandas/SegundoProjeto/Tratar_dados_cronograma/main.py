import pandas as pd
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# Carregar o arquivo Excel
df = pd.read_excel('Cronograma de processos - 2025 (Responses).xlsx')
dia_hora = datetime.now().strftime('%Y%m%d_%H%M%S')
caminho_Salvar = r"//apolo/Governanca/ReportBI/BaseStatusReportCronogram"

# Criar um novo dataframe para armazenar os dados tratados
data = {
    'Processo': [],
    'Justificativa': [],
    'Descrição': [],
    'Data de entrega': [],
    'Diretoria': [],
    'Setores': [],
    'Endereço de e-mail': [],
    'Carimbo de data/hora': []
}

# Iterar sobre as linhas do dataframe original
for index, row in df.iterrows():
    for i in range(0, 5):  # Existem 5 processos no arquivo
        processo_col = f'Processo {i+1}'
        if pd.notna(row[processo_col]):
            data['Processo'].append(row[processo_col])   
            if i == 0:
                desloc = [4, 3, 1]
            else:
                desloc = [1, 2, 3]
                
            pos_inicial = df.columns.get_loc(processo_col)
            justificativa_col = df.columns[pos_inicial + desloc[0]] if pos_inicial + desloc[0] < len(df.columns) else None
            descricao_col = df.columns[pos_inicial + desloc[1]] if pos_inicial + desloc[1] < len(df.columns) else None
            data_entrega_col = df.columns[pos_inicial + desloc[2]] if pos_inicial + desloc[2] < len(df.columns) else None
            
            data['Justificativa'].append(row[justificativa_col] if justificativa_col else '')
            data['Descrição'].append(row[descricao_col] if descricao_col else '')
            data['Data de entrega'].append(f"{row[data_entrega_col]}/{row['Carimbo de data/hora'].year}" if data_entrega_col and pd.notna(row[data_entrega_col]) and isinstance(row['Carimbo de data/hora'], datetime) else '')
            data['Diretoria'].append(row['Diretoria'])
            data['Setores'].append(row[f"Setores {row['Diretoria']}"])
            data['Endereço de e-mail'].append(row['Endereço de e-mail'])
            data['Carimbo de data/hora'].append(row['Carimbo de data/hora'])

# Criar um novo dataframe com os dados tratados
df_tratado = pd.DataFrame(data)

# Converter a coluna Data de entrega para formato ISO 8601
df_tratado['Data de entrega'] = pd.to_datetime(df_tratado['Data de entrega'], format="%d/%m/%Y", errors='coerce').dt.strftime('%Y-%m-%dT%H:%M:%SZ')

# Criar o arquivo Excel e formatar como tabela
arquivo_excel = f"Resultado\\Cronograma_de_processos_tratado_{dia_hora}.xlsx"

try:
    # Salvar o DataFrame como Excel usando pandas
    df_tratado.to_excel(arquivo_excel, index=False)

    # Abrir o arquivo salvo com openpyxl
    wb = load_workbook(arquivo_excel)
    ws = wb.active

    # Criar a tabela no intervalo correspondente ao DataFrame
    ref = f"A1:{chr(64 + len(df_tratado.columns))}{len(df_tratado) + 1}"
    tabela = Table(displayName="TabelaProcessos", ref=ref)

    # Adicionar a tabela à planilha
    ws.add_table(tabela)

    # Salvar o arquivo Excel formatado
    wb.save(arquivo_excel)

    print(f"Dados tratados e salvos como tabela em '{arquivo_excel}'")
except OSError:
    print("Não foi possível salvar o arquivo, caminho não encontrado")
    exit(2)
except RuntimeError:
    print("Não foi possível finalizar a operação, verifique o arquivo de input")
    exit(5)
